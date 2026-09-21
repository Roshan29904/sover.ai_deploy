from pathlib import Path

from docx import Document
from langchain_core.tools import tool


@tool
def create_docx(title: str, content: str, output_path: str) -> str:
    """
    Create a DOCX document locally with a title and text content.
    """

    try:
        path = Path(output_path)

        path.parent.mkdir(parents=True, exist_ok=True)

        document = Document()

        document.add_heading(title, level=1)

        for paragraph in content.split("\n"):
            if paragraph.strip():
                document.add_paragraph(paragraph)

        document.save(path)

        return f"DOCUMENT_CREATED:{path.resolve()}"

    except Exception as e:
        return f"Unable to create DOCX: {e}"
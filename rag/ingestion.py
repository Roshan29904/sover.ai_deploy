from pathlib import Path

from langchain_community.document_loaders import (
    TextLoader,
    PyPDFLoader,
    Docx2txtLoader,
    UnstructuredPowerPointLoader,
)
from langchain_core.documents import Document
from langchain_text_splitters import RecursiveCharacterTextSplitter

from openpyxl import load_workbook


def load_document(file_path: str):

    path = Path(file_path)

    if not path.exists():
        raise FileNotFoundError(f"File not found: {file_path}")

    extension = path.suffix.lower()

    if extension == ".txt":
        loader = TextLoader(str(path), encoding="utf-8")
        documents = loader.load()

    elif extension == ".pdf":
        loader = PyPDFLoader(str(path))
        documents = loader.load()

    elif extension == ".docx":
        loader = Docx2txtLoader(str(path))
        documents = loader.load()

    elif extension == ".pptx":
        loader = UnstructuredPowerPointLoader(str(path))
        documents = loader.load()

    elif extension == ".xlsx":
        documents = load_excel(path)

    else:
        raise ValueError(f"Unsupported file type: {extension}")

    return documents


def load_excel(path: Path):

    workbook = load_workbook(path, data_only=True)

    documents = []

    for sheet in workbook.worksheets:

        rows = []

        for row in sheet.iter_rows(values_only=True):

            values = [
                str(value) for value in row
                if value is not None
            ]

            if values:
                rows.append(" | ".join(values))

        if rows:

            content = "\n".join(rows)

            documents.append(
                Document(
                    page_content=content,
                    metadata={
                        "source": str(path),
                        "sheet": sheet.title
                    }
                )
            )

    return documents


def load_and_split(file_path: str):

    documents = load_document(file_path)

    splitter = RecursiveCharacterTextSplitter(
        chunk_size=500,
        chunk_overlap=50
    )

    return splitter.split_documents(documents)
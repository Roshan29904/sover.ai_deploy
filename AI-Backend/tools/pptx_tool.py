from pathlib import Path

from pptx import Presentation
from langchain_core.tools import tool


@tool
def create_pptx(title: str, content: str, output_path: str) -> str:
    """
    Create a PowerPoint presentation locally.
    Each line in content becomes a bullet point on the slide.
    """

    try:
        path = Path(output_path)
        path.parent.mkdir(parents=True, exist_ok=True)

        presentation = Presentation()

        slide = presentation.slides.add_slide(
            presentation.slide_layouts[1]
        )

        slide.shapes.title.text = title

        text_frame = slide.placeholders[1].text_frame
        text_frame.clear()

        for line in content.strip().split("\n"):
            if line.strip():
                paragraph = text_frame.add_paragraph()
                paragraph.text = line.strip()
                paragraph.level = 0

        presentation.save(path)

        return f"DOCUMENT_CREATED:{path.resolve()}"

    except Exception as e:
        return f"Unable to create PowerPoint: {e}"